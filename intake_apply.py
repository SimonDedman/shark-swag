#!/usr/bin/env python3
"""Read a reviewed workbook back into items.json, the picks, and the document.

Refuses to guess. An accepted row missing a section, a subsection, an item name,
or a description stops the run and is named, because a wrong entry in the shared
document costs more than a second pass over the sheet.
"""
import re
import subprocess
import sys
from pathlib import Path

import openpyxl

import entries
import intake_classify
import intake_queue
import picks as picks_mod
from intake_sheet import COLS as CONTACT_SHEET_COLS, ranked_candidates
from paths import INTAKE, ROOT, SHEETS

VENV_PYTHON = "/home/simon/.venvs/scrape/bin/python"
REQUIRED = ["Section", "Subsection", "Item", "Description"]

# An index typed into Image pick is a position on the contact sheet, and
# nothing else. Anything that is not a bare non-negative integer is refused
# rather than skipped: a skipped fragment leaves picks.json holding an empty
# spec list, which picks.py defines as "looked at, no usable image exists" —
# a durable false statement bought with one stray hyphen.
PICK_INDEX = re.compile(r"[0-9]+")


def newest_workbook():
    books = sorted(INTAKE.glob("review_*.xlsx"))
    if not books:
        raise SystemExit(f"no review_*.xlsx in {INTAKE}")
    return books[-1]


def read_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Review"]
    headers = [c.value for c in ws[1]]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, "") for v in values):
            continue
        row = {h: ("" if v is None else v) for h, v in zip(headers, values)}
        row["Status"] = (str(row.get("Status") or "").strip().upper() or "HOLD")
        rows.append(row)
    return rows


def _pick_fragments(raw, qid, errors):
    """Split an Image pick cell into indices, complaining about anything else."""
    out = []
    for part in str(raw).split(","):
        part = part.strip()
        if PICK_INDEX.fullmatch(part):
            out.append(int(part))
        else:
            errors.append(f"{qid}: Image pick {part!r} is not a whole number; "
                          f"use the indices printed on the contact sheet, "
                          f"comma separated")
    return out


def validate(rows, queue=None):
    if queue is None:
        queue = {r["queue_id"]: r for r in intake_queue.load_queue()}
    errors = []
    for row in rows:
        if row["Status"] != "ACCEPT":
            continue
        raw_qid = row.get("Queue id")
        if not raw_qid:
            errors.append(
                "?: this row is ACCEPT but its Queue id is blank, so it cannot "
                "be traced to any submission. That is what a row typed by hand "
                "into the workbook looks like; give it a real queue id, or add "
                "it through the intake process instead.")
            continue
        qid = raw_qid
        if qid not in queue:
            errors.append(
                f"{qid}: this row is ACCEPT but no longer in the queue, which is "
                f"what applying it removes. It looks as though this workbook has "
                f"already been applied; applying it again would add every "
                f"accepted entry a second time under a second id.")
            continue
        for field in REQUIRED:
            if not str(row.get(field, "")).strip():
                errors.append(f"{qid}: {field} is empty on an ACCEPT row")
        section, sub = str(row.get("Section", "")), str(row.get("Subsection", ""))
        if section and sub:
            valid = intake_classify.valid_subs(section)
            if sub not in valid:
                errors.append(
                    f"{qid}: subsection {sub!r} does not exist in section {section}; "
                    f"valid names are {valid}")
        raw = str(row.get("Image pick", "")).strip()
        if raw:
            cands = ranked_candidates(queue.get(qid) or {})
            for index in _pick_fragments(raw, qid, errors):
                if index >= CONTACT_SHEET_COLS:
                    errors.append(
                        f"{qid}: Image pick {index} is off the "
                        f"{CONTACT_SHEET_COLS}-image contact sheet")
                elif index >= len(cands):
                    errors.append(
                        f"{qid}: Image pick {index} is beyond this row's "
                        f"{len(cands)} candidate(s)")
    return errors


def apply_row(row, data):
    """Insert one accepted row into `data` and return the id it was given."""
    new_id = entries.next_id(data)
    item = {
        "id": new_id,
        "item": str(row["Item"]).strip(),
        "age": str(row.get("Age", "")).strip(),
        "link": str(row["Link"]).strip(),
        "price": str(row.get("Approx. price", "")).strip(),
        "desc": str(row["Description"]).strip(),
        "by": str(row.get("Suggested by", "")).strip(),
        "ngo": str(row.get("NGO / proceeds to conservation?", "")).strip(),
    }
    for section in data["sections"]:
        if section["number"] != str(row["Section"]):
            continue
        for sub in section["subsections"]:
            if sub["name"] == str(row["Subsection"]):
                sub["items"].append(item)
                return new_id
    raise SystemExit(f"section {row['Section']} subsection {row['Subsection']} not found")


def _picked_specs(row, queue_row):
    """Turn "0,3" plus the row's candidates into picks.json specs.

    The index means a position on the contact sheet, so it is resolved through
    the same ranked_candidates() the sheet laid out. validate() has already
    refused anything malformed by the time main() gets here; the same refusal
    is repeated rather than assumed, because this is the one function that
    turns a number Simon typed into a URL.
    """
    raw = str(row.get("Image pick", "")).strip()
    if not raw or not queue_row:
        return []
    qid = row.get("Queue id") or "?"
    errors = []
    indices = _pick_fragments(raw, qid, errors)
    if errors:
        raise ValueError("; ".join(errors))
    cands = ranked_candidates(queue_row)
    specs = []
    for index in indices:
        if index >= len(cands):
            raise ValueError(f"{qid}: Image pick {index} is beyond this row's "
                             f"{len(cands)} candidate(s)")
        specs.append("U:" + cands[index]["url"])
    return specs


def _finalize_queue(accepted, rejected, queue_path=None):
    """Drop only the rows this run actually decided on: an ACCEPT that was
    applied, or a REJECT that was recorded. Everything else in queue.json
    survives untouched, including a submission the intake daemon staged in
    while this workbook was being filled in: a row's mere absence from the
    workbook is not a decision about it, so it must not be deleted.

    `queue_path` lets a caller (a test) point this at a temp file instead of
    the real intake/queue.json; production calls leave it as None and get
    intake_queue's own default.

    Also removes each finalized row's pending contact sheet
    (images/sheets/pending_<queue_id>.jpg), if one exists. next_queue_id maxes
    over the current queue only, so once a row leaves it a later submission
    can be handed the same id back; if that id's old sheet were left behind,
    write_workbook would link it to the new row and offer a contact sheet
    showing a different item's images. A held row's sheet is left alone: its
    id is still live and the sheet still belongs to it.
    """
    kwargs = {} if queue_path is None else {"path": queue_path}
    processed = {qid for qid, _entry_id in accepted} | set(rejected)
    rows = intake_queue.load_queue(**kwargs)
    intake_queue.save_queue([r for r in rows if r["queue_id"] not in processed], **kwargs)
    for qid in processed:
        (SHEETS / f"pending_{qid}.jpg").unlink(missing_ok=True)


def main(path=None, rebuild=True):
    path = Path(path) if path else newest_workbook()
    rows = read_workbook(path)
    queue = {r["queue_id"]: r for r in intake_queue.load_queue()}
    errors = validate(rows, queue)
    if errors:
        print(f"{len(errors)} problem(s); nothing was applied:")
        for e in errors:
            print("   ", e)
        raise SystemExit(1)

    data = entries.load_items()
    all_picks = picks_mod.load_picks()

    accepted, rejected, held = [], [], []
    for row in rows:
        qid = row.get("Queue id") or ""
        if row["Status"] == "ACCEPT":
            entry_id = apply_row(row, data)
            specs = _picked_specs(row, queue.get(qid))
            all_picks[entry_id] = dict(specs=specs, note=f"from intake {qid}")
            accepted.append((qid, entry_id))
        elif row["Status"] == "REJECT":
            queue_row = queue.get(qid, {})
            if queue_row.get("url_norm"):
                intake_queue.reject(queue_row["url_norm"], "rejected in review")
            rejected.append(qid)
        else:
            held.append(qid)

    entries.save_items(data)
    picks_mod.save_picks(all_picks)
    _finalize_queue(accepted, rejected)

    for _qid, entry_id in accepted:
        written, failures = picks_mod.render(entry_id, all_picks[entry_id]["specs"])
        for f in failures:
            print("   image failed:", f)

    print(f"{len(accepted)} accepted, {len(rejected)} rejected, {len(held)} held")
    for qid, entry_id in accepted:
        print(f"   {qid} -> {entry_id}")

    if rebuild and accepted:
        subprocess.run([VENV_PYTHON, str(ROOT / "build_docx.py")], check=True)
        subprocess.run([VENV_PYTHON, str(ROOT / "add_images.py")], check=True)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else None)
