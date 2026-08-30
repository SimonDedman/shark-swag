import json

import pytest

import intake_apply
import intake_queue
import intake_sheet


def sheet_row(**over):
    base = dict(queue_id="q9001", message_id="m1", received="2026-08-28T09:14:02-07:00",
                sender="Simon Dedman", note="", url="https://shop.com/products/socks",
                url_norm="https://shop.com/products/socks", fetch_status="OK", http=200,
                fetch_note="fetched", title="Shark socks", price="~£12.50",
                price_source="itemprop", ngo_text="5% donated.", section="1",
                sub="Adult", age="Adult", keyword="socks", duplicate_of="",
                candidates=[], picked="")
    base.update(over)
    return base


@pytest.fixture
def staged():
    """The queue as it stands when a freshly generated workbook is applied:
    every row in the workbook is still staged. validate() reads it to tell a
    first run from a second one."""
    def build(*queue_ids):
        return {qid: dict(sheet_row(queue_id=qid)) for qid in (queue_ids or ("q9001",))}
    return build


@pytest.fixture
def workbook(tmp_path):
    def build(rows, statuses, descriptions=None):
        path = tmp_path / "review.xlsx"
        intake_sheet.write_workbook(rows, path)
        import openpyxl
        wb = openpyxl.load_workbook(path)
        ws = wb["Review"]
        headers = [c.value for c in ws[1]]
        for i, (status, desc) in enumerate(zip(statuses, descriptions or [""] * len(rows)), 2):
            ws.cell(row=i, column=headers.index("Status") + 1, value=status)
            ws.cell(row=i, column=headers.index("Description") + 1, value=desc)
        wb.save(path)
        return path
    return build


def test_read_workbook_returns_one_dict_per_row(workbook):
    path = workbook([sheet_row()], ["ACCEPT"], ["Bamboo socks."])
    rows = intake_apply.read_workbook(path)
    assert len(rows) == 1
    assert rows[0]["Status"] == "ACCEPT"
    assert rows[0]["Description"] == "Bamboo socks."
    assert rows[0]["Queue id"] == "q9001"


def test_a_blank_status_reads_as_hold(workbook):
    path = workbook([sheet_row()], [None])
    assert intake_apply.read_workbook(path)[0]["Status"] == "HOLD"


def test_validate_refuses_an_accepted_row_with_no_description(workbook, staged):
    path = workbook([sheet_row()], ["ACCEPT"], [""])
    errors = intake_apply.validate(intake_apply.read_workbook(path), staged())
    assert errors
    assert "q9001" in errors[0]
    assert "Description" in errors[0]


def test_validate_refuses_an_accepted_row_with_no_section(workbook, staged):
    path = workbook([sheet_row(section="")], ["ACCEPT"], ["Bamboo socks."])
    errors = intake_apply.validate(intake_apply.read_workbook(path), staged())
    assert any("Section" in e for e in errors)


def test_validate_refuses_a_subsection_that_does_not_exist(workbook, staged):
    path = workbook([sheet_row(sub="Grown-ups")], ["ACCEPT"], ["Bamboo socks."])
    errors = intake_apply.validate(intake_apply.read_workbook(path), staged())
    assert any("Grown-ups" in e for e in errors)


def test_validate_passes_a_complete_row(workbook, staged):
    path = workbook([sheet_row()], ["ACCEPT"], ["Bamboo socks."])
    assert intake_apply.validate(intake_apply.read_workbook(path), staged()) == []


def test_rejected_and_held_rows_are_not_validated(workbook, staged):
    path = workbook([sheet_row(section=""), sheet_row(queue_id="q9002", section="")],
                    ["REJECT", "HOLD"], ["", ""])
    assert intake_apply.validate(intake_apply.read_workbook(path), staged("q9001", "q9002")) == []


def test_apply_row_inserts_into_the_named_subsection_with_a_fresh_id():
    data = {"sections": [{"number": "1", "title": "t", "subsections":
                          [{"name": "Adult", "items": [{"id": "e01", "item": "x",
                                                        "link": "https://a"}]}]}]}
    row = {"Status": "ACCEPT", "Section": "1", "Subsection": "Adult",
           "Item": "Shark socks", "Age": "Adult", "Link": "https://shop.com/p",
           "Approx. price": "~£12.50", "Description": "Bamboo socks.",
           "Suggested by": "Simon Dedman",
           "NGO / proceeds to conservation?": "5% donated.", "Image pick": "",
           "Queue id": "q9001"}
    new_id = intake_apply.apply_row(row, data)
    items = data["sections"][0]["subsections"][0]["items"]
    assert len(items) == 2
    assert items[1]["id"] == new_id
    assert items[1]["item"] == "Shark socks"
    assert items[1]["desc"] == "Bamboo socks."
    assert items[1]["by"] == "Simon Dedman"
    assert new_id != "e01"


def test_apply_row_does_not_disturb_existing_ids():
    data = {"sections": [{"number": "1", "title": "t", "subsections":
                          [{"name": "Adult", "items": [{"id": "e01", "item": "x",
                                                        "link": "https://a"}]}]}]}
    row = {"Status": "ACCEPT", "Section": "1", "Subsection": "Adult",
           "Item": "Shark socks", "Age": "Adult", "Link": "https://shop.com/p",
           "Approx. price": "", "Description": "d", "Suggested by": "s",
           "NGO / proceeds to conservation?": "", "Image pick": "", "Queue id": "q9001"}
    intake_apply.apply_row(row, data)
    assert data["sections"][0]["subsections"][0]["items"][0]["id"] == "e01"


# --- Task 11 addition: an Image pick index must be one Simon actually saw on
# the contact sheet (which lays out only the first CONTACT_SHEET_COLS
# candidates), and must exist among this row's own candidates. ---

def test_validate_refuses_an_image_pick_off_the_contact_sheet(workbook):
    path = workbook([sheet_row()], ["ACCEPT"], ["Bamboo socks."])
    rows = intake_apply.read_workbook(path)
    rows[0]["Image pick"] = str(intake_apply.CONTACT_SHEET_COLS)
    queue = {"q9001": {"candidates": [{"url": f"https://x/{i}.jpg"} for i in range(20)]}}
    errors = intake_apply.validate(rows, queue)
    assert errors
    assert "q9001" in errors[0]
    assert str(intake_apply.CONTACT_SHEET_COLS) in errors[0]


def test_validate_refuses_an_image_pick_past_this_rows_own_candidates(workbook):
    path = workbook([sheet_row()], ["ACCEPT"], ["Bamboo socks."])
    rows = intake_apply.read_workbook(path)
    rows[0]["Image pick"] = "3"
    queue = {"q9001": {"candidates": [{"url": "https://x/0.jpg"}, {"url": "https://x/1.jpg"}]}}
    errors = intake_apply.validate(rows, queue)
    assert errors
    assert "q9001" in errors[0]
    assert "3" in errors[0]


def test_validate_accepts_an_in_range_image_pick(workbook):
    path = workbook([sheet_row()], ["ACCEPT"], ["Bamboo socks."])
    rows = intake_apply.read_workbook(path)
    rows[0]["Image pick"] = "1"
    queue = {"q9001": {"candidates": [{"url": "https://x/0.jpg"}, {"url": "https://x/1.jpg"}]}}
    assert intake_apply.validate(rows, queue) == []


# --- Task 11 fix (post-review): the final queue rewrite must not delete a
# submission the daemon staged in after the workbook was generated. ---

def test_a_submission_staged_after_the_workbook_was_written_is_not_deleted(tmp_path, workbook):
    # Two rows go into the workbook, as of the moment it was generated:
    # q9001 will be accepted, q9002 will be held.
    path = workbook([sheet_row(), sheet_row(queue_id="q9002")], ["ACCEPT", "HOLD"],
                    ["Bamboo socks.", ""])
    rows = intake_apply.read_workbook(path)
    accepted = [(rows[0]["Queue id"], "e99")]
    rejected = []

    # Sometime between the workbook being generated and this run applying it,
    # the intake daemon staged a third submission, q9003, straight into
    # queue.json. It was never in the workbook and is neither accepted nor
    # rejected by this run.
    queue_path = tmp_path / "queue.json"
    intake_queue.save_queue(
        [{"queue_id": "q9001"}, {"queue_id": "q9002"}, {"queue_id": "q9003"}],
        path=queue_path)

    intake_apply._finalize_queue(accepted, rejected, queue_path=queue_path)

    survivors = {r["queue_id"] for r in intake_queue.load_queue(path=queue_path)}
    assert survivors == {"q9002", "q9003"}


# --- Critical 2: the workbook stays in intake/ after a run and
# newest_workbook() picks it up again, so a second run — exactly what a failed
# build_docx.py invites — appended every accepted entry a second time under a
# second id. ---

def test_validate_refuses_a_workbook_that_has_already_been_applied(workbook, staged):
    """After a run the accepted row is gone from the queue, which is the only
    durable record that it was applied."""
    path = workbook([sheet_row()], ["ACCEPT"], ["Bamboo socks."])
    rows = intake_apply.read_workbook(path)
    errors = intake_apply.validate(rows, queue={})       # q9001 already applied
    assert errors
    assert "q9001" in errors[0]
    assert "already been applied" in errors[0]


def test_a_first_run_is_unaffected_by_the_already_applied_check(workbook, staged):
    path = workbook([sheet_row()], ["ACCEPT"], ["Bamboo socks."])
    assert intake_apply.validate(intake_apply.read_workbook(path), staged()) == []


def test_a_second_run_refuses_even_when_the_row_carried_no_image_pick(workbook):
    """The row with an image pick was protected only by accident, because its
    candidate list left with the queue row. A row without one sailed through."""
    path = workbook([sheet_row(picked="")], ["ACCEPT"], ["Bamboo socks."])
    rows = intake_apply.read_workbook(path)
    assert rows[0]["Image pick"] in (None, "")
    assert intake_apply.validate(rows, queue={})


def test_held_rows_do_not_trip_the_already_applied_check(workbook):
    """A HOLD row is still in the queue after a run, and an unrelated missing
    queue id must not be invented as an error for a row nobody accepted."""
    path = workbook([sheet_row()], ["HOLD"], [""])
    assert intake_apply.validate(intake_apply.read_workbook(path), queue={}) == []


# --- Important 6: '0-2', '0;1', '0.0', and '-1' all validated clean and
# produced no specs, so picks.json recorded {"specs": []}, which picks.py
# defines as "looked at, and there is no usable image". ---

MALFORMED = ["0-2", "0;1", "0.0", "-1"]


@pytest.mark.parametrize("bad", MALFORMED)
def test_validate_refuses_a_malformed_image_pick(workbook, staged, bad):
    path = workbook([sheet_row()], ["ACCEPT"], ["Bamboo socks."])
    rows = intake_apply.read_workbook(path)
    rows[0]["Image pick"] = bad
    queue = {"q9001": {"candidates": [{"url": "https://x/0.jpg", "src": "img"},
                                      {"url": "https://x/1.jpg", "src": "img"}]}}
    errors = intake_apply.validate(rows, queue)
    assert errors, f"{bad!r} passed validation"
    assert "q9001" in errors[0]
    assert any(bad in e or bad.split("-")[0] in e for e in errors)


@pytest.mark.parametrize("bad", MALFORMED)
def test_picked_specs_refuses_a_malformed_image_pick(bad):
    row = {"Image pick": bad, "Queue id": "q9001"}
    queue_row = {"candidates": [{"url": "https://x/0.jpg", "src": "img"}]}
    with pytest.raises(ValueError) as exc:
        intake_apply._picked_specs(row, queue_row)
    assert "q9001" in str(exc.value)


def test_a_malformed_pick_never_becomes_an_empty_spec_list(workbook, staged):
    """The whole point: an empty spec list is a documented statement that the
    entry has no usable image. A stray hyphen must not make that statement."""
    path = workbook([sheet_row()], ["ACCEPT"], ["Bamboo socks."])
    rows = intake_apply.read_workbook(path)
    rows[0]["Image pick"] = "0-2"
    queue = {"q9001": {"candidates": [{"url": "https://x/0.jpg", "src": "img"}]}}
    assert intake_apply.validate(rows, queue)
    with pytest.raises(ValueError):
        intake_apply._picked_specs(rows[0], queue["q9001"])


# --- Important 10: the three seams where the modules meet had no tests at all.
# _picked_specs is the one function that turns a number Simon typed into a URL,
# and the index-to-candidate contract was asserted nowhere, only its bounds. ---

def test_picked_specs_maps_an_index_to_that_candidates_url():
    cands = [{"url": f"https://x/{i}.jpg", "title": "", "src": "img"} for i in range(3)]
    row = {"Image pick": "1", "Queue id": "q9001"}
    specs = intake_apply._picked_specs(row, {"candidates": cands, "title": ""})
    assert specs == ["U:" + cands[1]["url"]]


def test_picked_specs_reads_the_index_off_the_ranked_order_the_sheet_drew():
    """The sheet numbers cells by rank, so the number Simon types means a rank,
    not a position in the harvest. Resolving it against the raw list would hand
    him a different picture from the one he chose."""
    cands = [{"url": "https://x/furniture.png", "title": "", "src": "img"},
             {"url": "https://x/product.jpg", "title": "", "src": "shopify-product"}]
    row = {"Image pick": "0", "Queue id": "q9001"}
    specs = intake_apply._picked_specs(row, {"candidates": cands, "title": ""})
    assert specs == ["U:https://x/product.jpg"]


def test_picked_specs_takes_several_indices_in_order():
    cands = [{"url": f"https://x/{i}.jpg", "title": "", "src": "img"} for i in range(4)]
    row = {"Image pick": "2, 0", "Queue id": "q9001"}
    specs = intake_apply._picked_specs(row, {"candidates": cands, "title": ""})
    assert specs == ["U:https://x/2.jpg", "U:https://x/0.jpg"]


def test_picked_specs_is_empty_when_no_image_was_chosen():
    assert intake_apply._picked_specs({"Image pick": "", "Queue id": "q9001"},
                                      {"candidates": []}) == []


# --- Important 10: main() had no test at all, and it is where every other
# module's output meets items.json. Runs entirely against temp state: no
# network, no docx rebuild, and the real items.json is never opened. ---

ITEMS_SEED = {
    "sections": [
        {"number": "1", "title": "Basics", "subsections": [
            {"name": "Adult", "items": [{"id": "e01", "item": "x", "link": "https://a",
                                         "age": "Adult", "price": "", "desc": "d",
                                         "by": "s", "ngo": ""}]},
        ]},
    ]
}


@pytest.fixture
def isolated_apply(tmp_path, monkeypatch):
    """items.json, existing.json, picks.json, and the whole intake/ directory
    all point at tmp_path. Nothing here touches the real files."""
    items = tmp_path / "items.json"
    items.write_text(json.dumps(ITEMS_SEED))
    existing = tmp_path / "existing.json"
    existing.write_text("[]")
    picks = tmp_path / "picks.json"
    picks.write_text("{}")

    import entries
    import picks as picks_mod
    monkeypatch.setattr(entries, "ITEMS", items)
    monkeypatch.setattr(entries, "EXISTING_JSON", existing)
    monkeypatch.setattr(entries, "PICKS_JSON", picks)
    monkeypatch.setattr(picks_mod, "PICKS_JSON", picks)
    monkeypatch.setattr(intake_queue, "QUEUE", tmp_path / "queue.json")
    monkeypatch.setattr(intake_queue, "REJECTED", tmp_path / "rejected.json")
    monkeypatch.setattr(intake_queue, "LEDGER", tmp_path / "ledger.json")
    monkeypatch.setattr(picks_mod, "render", lambda entry_id, specs: ([], []))
    # _finalize_queue deletes a processed row's contact sheet; point it at
    # tmp_path too, so a real queue id never causes main() to touch (or
    # delete from) the project's actual images/sheets directory.
    monkeypatch.setattr(intake_apply, "SHEETS", tmp_path / "sheets")
    return tmp_path


def test_main_applies_one_accepted_row_end_to_end(isolated_apply, workbook, capsys):
    import entries
    import picks as picks_mod
    row = sheet_row(candidates=[{"url": "https://x/0.jpg", "title": "", "src": "img"},
                                {"url": "https://x/1.jpg", "title": "", "src": "img"}])
    intake_queue.save_queue([row])
    path = workbook([row], ["ACCEPT"], ["Bamboo socks."])
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb["Review"]
    headers = [c.value for c in ws[1]]
    ws.cell(row=2, column=headers.index("Image pick") + 1, value="1")
    wb.save(path)

    intake_apply.main(path=path, rebuild=False)

    data = entries.load_items()
    items = data["sections"][0]["subsections"][0]["items"]
    assert [i["item"] for i in items] == ["x", "Shark socks"]
    new_id = items[1]["id"]
    assert new_id != "e01"
    assert items[1]["desc"] == "Bamboo socks."

    assert picks_mod.load_picks()[new_id]["specs"] == ["U:https://x/1.jpg"]
    assert picks_mod.load_picks()[new_id]["note"] == "from intake q9001"
    assert intake_queue.load_queue() == []
    assert "1 accepted, 0 rejected, 0 held" in capsys.readouterr().out


def test_main_refuses_the_same_workbook_a_second_time(isolated_apply, workbook, capsys):
    import entries
    row = sheet_row()
    intake_queue.save_queue([row])
    path = workbook([row], ["ACCEPT"], ["Bamboo socks."])

    intake_apply.main(path=path, rebuild=False)
    before = entries.load_items()

    with pytest.raises(SystemExit):
        intake_apply.main(path=path, rebuild=False)

    assert entries.load_items() == before, "a second run duplicated the entry"
    assert "already been applied" in capsys.readouterr().out


def test_main_records_a_rejection_and_leaves_items_alone(isolated_apply, workbook):
    import entries
    row = sheet_row(url_norm="https://shop.com/products/socks")
    intake_queue.save_queue([row])
    path = workbook([row], ["REJECT"], [""])
    before = entries.load_items()

    intake_apply.main(path=path, rebuild=False)

    assert entries.load_items() == before
    assert "https://shop.com/products/socks" in intake_queue.load_rejected()
    assert intake_queue.load_queue() == []


def test_main_leaves_a_held_row_in_the_queue(isolated_apply, workbook):
    row = sheet_row()
    intake_queue.save_queue([row])
    path = workbook([row], ["HOLD"], [""])
    intake_apply.main(path=path, rebuild=False)
    assert [r["queue_id"] for r in intake_queue.load_queue()] == ["q9001"]


# --- New feature: the owner must be told, in the run's own output, which
# accepted rows went in with Claude's draft untouched - not refused, since a
# draft may be correct as it stands, but never silent either. ---

def test_main_reports_a_row_that_kept_the_draft_verbatim(isolated_apply, workbook, capsys):
    row = sheet_row(description_draft="Bamboo socks with a hammerhead motif.")
    intake_queue.save_queue([row])
    path = workbook([row], ["ACCEPT"], ["Bamboo socks with a hammerhead motif."])

    intake_apply.main(path=path, rebuild=False)

    assert "q9001 kept the drafted description verbatim" in capsys.readouterr().out


def test_main_says_nothing_once_the_draft_has_been_edited(isolated_apply, workbook, capsys):
    row = sheet_row(description_draft="Bamboo socks with a hammerhead motif.")
    intake_queue.save_queue([row])
    path = workbook([row], ["ACCEPT"],
                    ["Bamboo socks, hand-knitted, with a hammerhead motif."])

    intake_apply.main(path=path, rebuild=False)

    assert "kept the drafted description verbatim" not in capsys.readouterr().out


def test_main_says_nothing_when_there_was_never_a_draft(isolated_apply, workbook, capsys):
    row = sheet_row(description_draft="")
    intake_queue.save_queue([row])
    path = workbook([row], ["ACCEPT"], ["Bamboo socks with a hammerhead motif."])

    intake_apply.main(path=path, rebuild=False)

    assert "kept the drafted description verbatim" not in capsys.readouterr().out


def test_main_applies_nothing_when_any_row_fails_validation(isolated_apply, workbook):
    import entries
    rows = [sheet_row(), sheet_row(queue_id="q9002")]
    intake_queue.save_queue(rows)
    path = workbook(rows, ["ACCEPT", "ACCEPT"], ["Bamboo socks.", ""])
    before = entries.load_items()
    with pytest.raises(SystemExit):
        intake_apply.main(path=path, rebuild=False)
    assert entries.load_items() == before
    assert len(intake_queue.load_queue()) == 2


# --- Residual fix: a blank Queue id was refused with the "already applied"
# message, which sends the reader hunting for the wrong problem. ---

def test_validate_gives_a_hand_typed_row_its_own_message(workbook):
    """A row typed straight into the workbook carries no Queue id at all,
    which is a different problem from a row whose id fell out of the queue
    because the workbook was already applied, and needs its own words."""
    path = workbook([sheet_row(queue_id="")], ["ACCEPT"], ["Bamboo socks."])
    rows = intake_apply.read_workbook(path)
    assert rows[0].get("Queue id") in (None, "")
    errors = intake_apply.validate(rows, queue={"q9001": sheet_row()})
    assert errors
    assert "already been applied" not in errors[0]
    assert "no Queue id" in errors[0] or "Queue id is blank" in errors[0]


# --- Residual fix: next_queue_id hands out ids from the *current* queue, so a
# processed row's id can be recycled onto a brand-new submission. If the old
# row's contact sheet is still on disk under that id, write_workbook offers it
# to the new row, which may have found no images of its own at all. ---

def test_finalize_queue_deletes_the_sheet_for_an_accepted_row(tmp_path, monkeypatch):
    monkeypatch.setattr(intake_apply, "SHEETS", tmp_path)
    (tmp_path / "pending_q9001.jpg").write_bytes(b"x")
    queue_path = tmp_path / "queue.json"
    intake_queue.save_queue([{"queue_id": "q9001"}], path=queue_path)

    intake_apply._finalize_queue([("q9001", "e99")], [], queue_path=queue_path)

    assert not (tmp_path / "pending_q9001.jpg").exists()


def test_finalize_queue_deletes_the_sheet_for_a_rejected_row(tmp_path, monkeypatch):
    monkeypatch.setattr(intake_apply, "SHEETS", tmp_path)
    (tmp_path / "pending_q9002.jpg").write_bytes(b"x")
    queue_path = tmp_path / "queue.json"
    intake_queue.save_queue([{"queue_id": "q9002"}], path=queue_path)

    intake_apply._finalize_queue([], ["q9002"], queue_path=queue_path)

    assert not (tmp_path / "pending_q9002.jpg").exists()


def test_finalize_queue_leaves_a_held_rows_sheet_alone(tmp_path, monkeypatch):
    monkeypatch.setattr(intake_apply, "SHEETS", tmp_path)
    (tmp_path / "pending_q9001.jpg").write_bytes(b"x")
    (tmp_path / "pending_q9003.jpg").write_bytes(b"x")
    queue_path = tmp_path / "queue.json"
    intake_queue.save_queue([{"queue_id": "q9001"}, {"queue_id": "q9003"}], path=queue_path)

    # q9003 is held: neither accepted nor rejected, so it stays in the queue
    # and its sheet must survive.
    intake_apply._finalize_queue([("q9001", "e99")], [], queue_path=queue_path)

    assert not (tmp_path / "pending_q9001.jpg").exists()
    assert (tmp_path / "pending_q9003.jpg").exists()


def test_finalize_queue_is_not_an_error_when_the_sheet_is_already_gone(tmp_path, monkeypatch):
    monkeypatch.setattr(intake_apply, "SHEETS", tmp_path)
    queue_path = tmp_path / "queue.json"
    intake_queue.save_queue([{"queue_id": "q9001"}], path=queue_path)

    intake_apply._finalize_queue([("q9001", "e99")], [], queue_path=queue_path)  # no raise


# --- Workbook dropdowns: the Section cell now carries a descriptor
# ("5. \U0001f3a8 Arts, crafts & stationery"), typed by the dropdown built in
# intake_sheet, but a bare number typed by hand or left over from an older
# workbook must keep working, since validate() and apply_row compare it
# against items.json's plain section number. ---

def test_normalise_section_strips_a_full_descriptor():
    assert intake_apply.normalise_section("5. \U0001f3a8 Arts, crafts & stationery") == "5"


def test_normalise_section_accepts_a_bare_number_unchanged():
    assert intake_apply.normalise_section("5") == "5"


def test_normalise_section_leaves_a_blank_cell_blank():
    assert intake_apply.normalise_section("") == ""
    assert intake_apply.normalise_section(None) == ""


def test_read_workbook_normalises_a_dropdown_written_section(workbook):
    """A workbook built by intake_sheet.write_workbook now shows the Section
    cell as a full descriptor; read_workbook must hand validate() and
    apply_row a plain number."""
    path = workbook([sheet_row()], ["ACCEPT"], ["Bamboo socks."])
    rows = intake_apply.read_workbook(path)
    assert rows[0]["Section"] == "1"


def test_validate_still_passes_a_descriptor_shaped_section(workbook, staged):
    path = workbook([sheet_row()], ["ACCEPT"], ["Bamboo socks."])
    rows = intake_apply.read_workbook(path)
    # write_workbook already turned "1" into a descriptor; confirm the round
    # trip through read_workbook produced the plain number validate() expects.
    assert rows[0]["Section"] == "1"
    assert intake_apply.validate(rows, staged()) == []


def test_apply_row_accepts_a_hand_typed_bare_section_number():
    """apply_row itself is never handed a descriptor directly - read_workbook
    normalises first - but it must keep accepting the bare number a caller
    (or an older workbook) hands it."""
    data = {"sections": [{"number": "5", "title": "t", "subsections":
                          [{"name": "Stickers", "items": []}]}]}
    row = {"Status": "ACCEPT", "Section": "5", "Subsection": "Stickers",
           "Item": "Shark stickers", "Age": "Kids", "Link": "https://shop.com/p",
           "Approx. price": "£3", "Description": "A sheet of stickers.",
           "Suggested by": "Simon Dedman",
           "NGO / proceeds to conservation?": "", "Image pick": "", "Queue id": "q9001"}
    new_id = intake_apply.apply_row(row, data)
    assert data["sections"][0]["subsections"][0]["items"][0]["item"] == "Shark stickers"
    assert new_id


# --- Adding the NGO? column must not disturb any header intake_apply reads
# by name: it inserts before, and reads none of the new column itself. ---

HEADERS_READ_BY_NAME = ["Status", "Section", "Subsection", "Item", "Age", "Link",
                        "Approx. price", "Description", "Suggested by",
                        "NGO / proceeds to conservation?", "Image pick", "Queue id"]


def test_every_header_intake_apply_reads_by_name_still_exists_verbatim():
    headers = [h for h, *_ in intake_sheet.COLUMNS]
    for h in HEADERS_READ_BY_NAME:
        assert h in headers, f"{h!r} missing from intake_sheet.COLUMNS"


def test_ngo_flag_column_is_not_required_and_does_not_break_a_round_trip(workbook):
    """NGO? is a filtering aid only; an ACCEPT row must still apply cleanly
    whatever its NGO? cell says."""
    path = workbook([sheet_row()], ["ACCEPT"], ["Bamboo socks."])
    rows = intake_apply.read_workbook(path)
    assert "NGO?" in rows[0]
    assert intake_apply.validate(rows, {"q9001": sheet_row()}) == []
