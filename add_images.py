#!/usr/bin/env python3
"""Edit shark-swag-suggestions.docx in place:

  * every entry's "> <item>" line becomes Heading 3, so it appears in the
    navigation tree under its age-band Heading 2;
  * a row of product images is inserted after each entry's NGO line.

Edited in place rather than rebuilt from build_docx.py so the manual wording
changes already made in the Gaps section survive.
"""
import copy, re, sys

import docx
from docx.shared import Pt, Inches, Emu
from PIL import Image

from entries import iter_entries
from paths import FINAL, DOCX

ROW_H_IN = 1.34          # 8 lines of 10.5pt Calibri at 1.15 line spacing
MAX_ROW_W_IN = 6.4       # usable width inside the default 1in margins
GAP_IN = 0.08


def name_to_ids():
    """Item name to the ids that carry it, in id order.

    The docx is matched by name because that is what the heading line holds,
    but the images are found by id, so inserting an entry cannot shift them.
    """
    mapping = {}
    for e in sorted(iter_entries(), key=lambda e: int(e.id[1:])):
        mapping.setdefault(e.item, []).append(e.id)
    return mapping


def images_for(entry_id):
    return sorted(FINAL.glob(f"{entry_id}_*.jpg"))


def insert_after(paragraph, doc):
    new_p = copy.deepcopy(paragraph._p)
    for child in list(new_p):
        new_p.remove(child)
    paragraph._p.addnext(new_p)
    return docx.text.paragraph.Paragraph(new_p, paragraph._parent)


def main():
    mapping = name_to_ids()
    doc = docx.Document(DOCX)
    paras = doc.paragraphs

    # locate each entry: the "> name" paragraph and its NGO line
    entries, i = [], 0
    while i < len(paras):
        t = paras[i].text
        if t.startswith("▸ "):
            head = paras[i]
            j = i + 1
            ngo = None
            while j < len(paras) and not paras[j].text.startswith("▸ "):
                if paras[j].text.startswith("NGO / proceeds to conservation?"):
                    ngo = paras[j]
                    break
                j += 1
            entries.append((head, ngo))
            i = j
        else:
            i += 1
    print(f"{len(entries)} entries found in the docx")

    used, placed, missing, headings = set(), 0, [], 0
    for head, ngo in entries:
        raw = head.text[2:]
        nm = re.sub(r"\s*\[(already in doc|splurge)\]\s*$", "", raw).strip()

        head.style = doc.styles["Heading 3"]
        head.paragraph_format.space_before = Pt(10)
        head.paragraph_format.space_after = Pt(2)
        headings += 1

        cands = [i for i in mapping.get(nm, []) if i not in used]
        if not cands:
            missing.append(nm)
            continue
        entry_id = cands[0]
        used.add(entry_id)

        files = images_for(entry_id)
        if not files or ngo is None:
            continue

        sizes = [Image.open(f).size for f in files]
        widths = [ROW_H_IN * w / h for w, h in sizes]
        total = sum(widths) + GAP_IN * (len(files) - 1)
        scale = min(1.0, MAX_ROW_W_IN / total)

        p = insert_after(ngo, doc)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        for n, (f, w) in enumerate(zip(files, widths)):
            if n:
                p.add_run("  ")
            p.add_run().add_picture(str(f), width=Inches(w * scale),
                                    height=Inches(ROW_H_IN * scale))
        placed += 1

    doc.save(DOCX)
    print(f"{headings} entry headings promoted to Heading 3")
    print(f"{placed} image rows inserted")
    if missing:
        print("unmatched entry names:", missing)


if __name__ == "__main__":
    main()
