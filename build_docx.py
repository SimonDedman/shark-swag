#!/usr/bin/env python3
"""Build the Shark Swag review docx from items.json, mirroring the Google Doc structure.

Rebuilding here produces the document WITHOUT product images and with entry names
as plain paragraphs. Run add_images.py afterwards to promote each entry to
Heading 3 (so it shows in the navigation tree) and insert its image row; the
images themselves come from harvest.py -> pool.py -> picks.py.
"""
import json
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from paths import ROOT, DOCX
import entries

DOC_URL = "https://docs.google.com/document/d/13DY4w0BPnldO8Zj7Ztq53U872CuVjLob-wtmxo6iWbE/edit"

# Entries already present in the Google Doc, reproduced in place so coverage is
# visible. They live in existing.json so entries.py and build_docx.py cannot
# disagree about them.
EXISTING = {}
for _e in entries.load_existing():
    EXISTING.setdefault((_e["section"], _e["sub"]), []).append(_e)

BLUE = RGBColor(0x0B, 0x55, 0x94)
GREY = RGBColor(0x66, 0x66, 0x66)


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = docx.oxml.shared.OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = docx.oxml.shared.OxmlElement("w:r")
    props = docx.oxml.shared.OxmlElement("w:rPr")
    colour = docx.oxml.shared.OxmlElement("w:color")
    colour.set(qn("w:val"), "0B5594")
    underline = docx.oxml.shared.OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(colour)
    props.append(underline)
    run.append(props)
    t = docx.oxml.shared.OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)
    return link


def field(doc, label, value, url=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(f"{label}: ")
    run.bold = True
    if url:
        add_hyperlink(p, url, value)
    else:
        p.add_run(value)
    return p


def entry(doc, e, existing=False):
    head = doc.add_paragraph()
    head.paragraph_format.space_before = Pt(10)
    head.paragraph_format.space_after = Pt(2)
    tag = head.add_run("▸ ")
    tag.bold = True
    name = head.add_run(e["item"])
    name.bold = True
    if existing:
        flag = head.add_run("   [already in doc]")
        flag.italic = True
        flag.font.size = Pt(9)
        flag.font.color.rgb = GREY
    price = e.get("price", "")
    if "splurge" in price.lower():
        flag = head.add_run("   [splurge]")
        flag.italic = True
        flag.font.size = Pt(9)
        flag.font.color.rgb = GREY

    field(doc, "Item", e["item"])
    field(doc, "Age", e["age"])
    field(doc, "Link", e["link"], url=e["link"])
    field(doc, "Approx. price", price)
    field(doc, "Description", e["desc"])
    field(doc, "Suggested by", e.get("by", "Simon Dedman (web search)"))
    field(doc, "NGO / proceeds to conservation?", e.get("ngo", ""))


def main():
    data = json.loads((ROOT / "items.json").read_text())
    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    title = doc.add_heading("🦈 Shark Swag — draft entries for review 🦈", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    intro = doc.add_paragraph()
    intro.add_run("Not for pasting — read this first. ").bold = True
    intro.add_run(
        f"These are candidate entries researched on {data['meta']['generated']} to seed the shared "
        "Google Doc, which had four entries in it. Structure, age bands and the seven-line entry template "
        "mirror the live doc exactly, so anything you approve pastes straight in. Entries marked "
    )
    intro.add_run("[already in doc]").italic = True
    intro.add_run(
        " are the existing colleague submissions, reproduced in place so you can see coverage at a glance — "
        "do not paste those back. Delete this box and any rejected entries before transferring."
    )

    notes = doc.add_paragraph()
    notes.add_run("Method and caveats. ").bold = True
    notes.add_run(
        "Brief was: US-first pricing, strong preference for conservation and NGO proceeds without excluding "
        "good commercial items, mostly under $50 with a few flagged splurges, and all four seller types "
        "(NGO/aquarium shops, independent makers, mainstream retail, experiences). Every link was "
        "HTTP-checked on the build date; three dead links were replaced and a handful of sites that block "
        "scripted requests (LEGO, Etsy, Bookshop.org, scuba.com) were routed to equivalent pages that "
        "respond. Prices are approximate and were read from listings or retailer pages at that moment — "
        "treat them as indicative, not quotes, and expect the print-on-demand and small-shop ranges to move. "
        "Where a shop's own pages state no donation, the NGO line says so rather than guessing. Source doc: "
    )
    add_hyperlink(notes, DOC_URL, "the shared Google Doc")
    notes.add_run(".")

    total = 0
    for section in data["sections"]:
        doc.add_page_break()
        doc.add_heading(f"{section['number']}. {section['title']}", level=1)
        if section.get("note"):
            n = doc.add_paragraph()
            r = n.add_run(section["note"])
            r.italic = True
            r.font.color.rgb = GREY

        for sub in section["subsections"]:
            doc.add_heading(sub["name"], level=2)
            if sub.get("hint"):
                h = doc.add_paragraph()
                r = h.add_run(sub["hint"])
                r.italic = True
                r.font.color.rgb = GREY

            for e in EXISTING.get((section["number"], sub["name"]), []):
                entry(doc, e, existing=True)

            if not sub["items"]:
                p = doc.add_paragraph()
                r = p.add_run("Nothing found that met the brief — see the gaps note at the end.")
                r.italic = True
                r.font.color.rgb = GREY
            for e in sub["items"]:
                entry(doc, e)
                total += 1

    doc.add_page_break()
    doc.add_heading("Gaps and judgement calls", level=1)
    for line in [
        "Optics & tech (section 4) is empty. Waterproof cameras and binoculars are a general-purpose "
        "purchase with no shark angle and no conservation link, so anything added would have been "
        "padding. ",
        "Baby & toddler clothing is thinner than the other bands. NGO shops mostly start at kids' sizes; "
        "the three entries listed are the only ones found with either a genuine conservation tie or an "
        "aquarium behind them.",
        "Shark Angels' shop redirects to an update notice rather than a storefront, so it is listed with a "
        "warning rather than as a live recommendation.",
        "Sharkopoly and the We Love Sharks store are commercial: the site publishes shark content and sells "
        "merch, but I found no charity registration or stated donation, so it is not credited as an NGO.",
        "Redbubble and Bonfire stores (MISS, Sharks4Kids, AWSC) are print-on-demand, so the organisation's "
        "cut is a royalty rather than the full margin. Buying direct from an org's own shop sends more.",
        "Two book links are Amazon search results rather than product pages (Good Night Sharks, Shark Girl) "
        "because editions vary and I would rather not pin a specific ASIN that goes out of print.",
        "Everything priced in GBP or AUD is converted at rough current rates for the approximate USD figure; "
        "the seller charges in their own currency.",
    ]:
        p = doc.add_paragraph(line, style="List Bullet")
        p.paragraph_format.space_after = Pt(4)

    out = DOCX
    doc.save(out)
    print(f"{total} new entries + {sum(len(v) for v in EXISTING.values())} existing -> {out}")


if __name__ == "__main__":
    main()
